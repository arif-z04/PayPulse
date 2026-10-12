import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_checklist_docx(output_path):
    doc = docx.Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x33, 0x41, 0x55) # Slate 700

    # Colors
    PRIMARY = RGBColor(0x1E, 0x3A, 0x8A) # Blue 900
    SECONDARY = RGBColor(0x02, 0x84, 0xC7) # Sky 600
    DARK_TEXT = RGBColor(0x0F, 0x17, 0x2A) # Slate 900
    CHECK_COLOR = RGBColor(0x05, 0x96, 0x69) # Emerald 600

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("PayPulse — Database Operations Checklist")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("Authoritative Technical Checklist of Database Operations Implemented in PayPulse")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Metadata Table
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    meta = [
        ("Project Platform", "PayPulse Enterprise Payroll, Attendance & Leave Management"),
        ("Database Engine", "MySQL Community Server 8.4 LTS (InnoDB Storage Engine)"),
        ("Character Set & Collation", "utf8mb4 / utf8mb4_unicode_ci (Strict SQL Modes)"),
        ("Backend Framework & ORM", "ASP.NET Core (.NET 10) / Entity Framework Core 10 (Pomelo MySQL)"),
    ]

    for idx, (k, v) in enumerate(meta):
        row = table.rows[idx]
        cell_k, cell_v = row.cells[0], row.cells[1]
        cell_k.width = Inches(2.2)
        cell_v.width = Inches(4.3)
        
        pk = cell_k.paragraphs[0]
        pk.paragraph_format.space_after = Pt(2)
        pk.paragraph_format.space_before = Pt(2)
        rk = pk.add_run(k)
        rk.font.bold = True
        rk.font.size = Pt(9.5)
        rk.font.color.rgb = DARK_TEXT
        
        pv = cell_v.paragraphs[0]
        pv.paragraph_format.space_after = Pt(2)
        pv.paragraph_format.space_before = Pt(2)
        rv = pv.add_run(v)
        rv.font.size = Pt(9.5)

        # Light gray background for key
        shd_k = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
        cell_k._tc.get_or_add_tcPr().append(shd_k)
        shd_v = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
        cell_v._tc.get_or_add_tcPr().append(shd_v)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(title)
        run.font.name = 'Calibri'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = PRIMARY
        
        # Add a subtle border bottom line
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="2" w:color="CBD5E1"/></w:pBdr>')
        p._p.get_or_add_pPr().append(pBdr)

    def add_sub_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(title)
        run.font.name = 'Calibri'
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = SECONDARY

    def add_check_item(name, checked=True):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.25)
        
        run_box = p.add_run("☑  " if checked else "☐  ")
        run_box.font.bold = True
        run_box.font.size = Pt(11)
        run_box.font.color.rgb = CHECK_COLOR if checked else RGBColor(0x94, 0xA3, 0xB8)
        
        run_text = p.add_run(name)
        run_text.font.size = Pt(10)
        run_text.font.color.rgb = DARK_TEXT

    # 1. DDL
    add_section_header("1. Data Definition Language (DDL) Operations")
    
    add_sub_header("Database & Schema Initialization")
    add_check_item("Create Database (paypulse_db)")
    add_check_item("Configure Character Set (utf8mb4) & Collation (utf8mb4_unicode_ci)")
    add_check_item("Foreign Key Checks Toggling (FOREIGN_KEY_CHECKS = 0/1)")

    add_sub_header("Table Definitions (13 Core Relational Entities)")
    tables = [
        "Create Table: User_Account",
        "Create Table: Department",
        "Create Table: Designation",
        "Create Table: Employee",
        "Create Table: Salary_Structure",
        "Create Table: Attendance_Log",
        "Create Table: Leave_Type",
        "Create Table: Leave_Allocation",
        "Create Table: Leave_Request",
        "Create Table: Monthly_Payroll",
        "Create Table: Payroll_Component_Detail",
        "Create Table: Disbursement_Record",
        "Create Table: System_Audit_Log"
    ]
    for t in tables:
        add_check_item(t)

    add_sub_header("Primary Key Constraints (13 PKs)")
    pks = [
        "Primary Key: User_Account(User_ID)",
        "Primary Key: Department(Department_ID)",
        "Primary Key: Designation(Designation_ID)",
        "Primary Key: Employee(Employee_ID)",
        "Primary Key: Salary_Structure(Structure_ID)",
        "Primary Key: Attendance_Log(Log_ID)",
        "Primary Key: Leave_Type(Leave_Type_ID)",
        "Primary Key: Leave_Allocation(Allocation_ID)",
        "Primary Key: Leave_Request(Request_ID)",
        "Primary Key: Monthly_Payroll(Payroll_ID)",
        "Primary Key: Payroll_Component_Detail(Detail_ID)",
        "Primary Key: Disbursement_Record(Disbursement_ID)",
        "Primary Key: System_Audit_Log(Audit_ID)"
    ]
    for pk in pks:
        add_check_item(pk)

    add_sub_header("Foreign Key Referential Integrity Constraints (19 FKs)")
    fks = [
        "Foreign Key: fk_designation_department (ON DELETE RESTRICT)",
        "Foreign Key: fk_employee_department (ON DELETE RESTRICT)",
        "Foreign Key: fk_employee_designation (ON DELETE RESTRICT)",
        "Foreign Key: fk_employee_user (ON DELETE CASCADE)",
        "Foreign Key: fk_employee_manager (ON DELETE SET NULL)",
        "Foreign Key: fk_department_head (ON DELETE SET NULL)",
        "Foreign Key: fk_salary_structure_employee (ON DELETE CASCADE)",
        "Foreign Key: fk_attendance_employee (ON DELETE CASCADE)",
        "Foreign Key: fk_leave_allocation_employee (ON DELETE CASCADE)",
        "Foreign Key: fk_leave_allocation_type (ON DELETE RESTRICT)",
        "Foreign Key: fk_leave_request_employee (ON DELETE CASCADE)",
        "Foreign Key: fk_leave_request_type (ON DELETE RESTRICT)",
        "Foreign Key: fk_leave_request_approver (ON DELETE SET NULL)",
        "Foreign Key: fk_monthly_payroll_employee (ON DELETE RESTRICT)",
        "Foreign Key: fk_monthly_payroll_user (ON DELETE SET NULL)",
        "Foreign Key: fk_payroll_component_payroll (ON DELETE CASCADE)",
        "Foreign Key: fk_disbursement_payroll (ON DELETE RESTRICT)",
        "Foreign Key: fk_disbursement_user (ON DELETE SET NULL)",
        "Foreign Key: fk_audit_user (ON DELETE SET NULL)"
    ]
    for fk in fks:
        add_check_item(fk)

    add_sub_header("Candidate & Unique Key Constraints (14 Keys)")
    uqs = [
        "Unique Key: uq_user_username",
        "Unique Key: uq_dept_name",
        "Unique Key: uq_dept_code",
        "Composite Unique Key: uq_dept_designation_title (Department_ID, Title)",
        "Unique Key: uq_emp_code",
        "Unique Key: uq_emp_user_id (1:1 User-Employee Profile)",
        "Unique Key: uq_emp_email",
        "Unique Key: uq_emp_nid",
        "Composite Unique Key: uq_emp_work_date (Employee_ID, Work_Date)",
        "Unique Key: uq_leave_type_name",
        "Composite Unique Key: uq_emp_leave_year (Employee_ID, Leave_Type_ID, Calendar_Year)",
        "Composite Unique Key: uq_emp_payroll_period (Employee_ID, Payroll_Year, Payroll_Month)",
        "Unique Key: uq_disb_payroll_id (1:1 Settlement Link)",
        "Unique Key: uq_disb_txref (BEFTN Transaction Reference)"
    ]
    for uq in uqs:
        add_check_item(uq)

    add_sub_header("Domain CHECK Constraints (31 Rules)")
    checks = [
        "Check: chk_user_username_minlen",
        "Check: chk_dept_code_format",
        "Check: chk_desig_salary_min",
        "Check: chk_desig_salary_range",
        "Check: chk_emp_name_nonempty",
        "Check: chk_emp_age_adult",
        "Check: chk_salary_basic_positive",
        "Check: chk_salary_hra_nonnegative",
        "Check: chk_salary_med_nonnegative",
        "Check: chk_salary_conv_nonnegative",
        "Check: chk_salary_tax_rate_range",
        "Check: chk_salary_pf_rate_range",
        "Check: chk_salary_effective_dates",
        "Check: chk_att_hours_range",
        "Check: chk_att_ot_range",
        "Check: chk_att_checkout_after_checkin",
        "Check: chk_leave_type_days_range",
        "Check: chk_alloc_year_range",
        "Check: chk_alloc_allocated_positive",
        "Check: chk_alloc_used_nonnegative",
        "Check: chk_alloc_remaining_nonnegative",
        "Check: chk_alloc_balance_identity",
        "Check: chk_req_date_order",
        "Check: chk_req_total_days_positive",
        "Check: chk_pay_month_range",
        "Check: chk_pay_year_range",
        "Check: chk_pay_working_days_range",
        "Check: chk_pay_present_days_range",
        "Check: chk_pay_gross_geq_basic",
        "Check: chk_pay_net_nonnegative",
        "Check: chk_pay_math_consistency"
    ]
    for chk in checks:
        add_check_item(chk)

    add_sub_header("Performance B-Tree Indexes (16 Indexes)")
    indexes = [
        "Index: idx_employee_department_status",
        "Index: idx_employee_designation_status",
        "Index: idx_employee_manager",
        "Index: idx_employee_hire_date",
        "Index: idx_attendance_workdate_status",
        "Index: idx_attendance_emp_date_status",
        "Index: idx_salary_structure_emp_active",
        "Index: idx_leave_request_status_emp",
        "Index: idx_leave_request_dates_status",
        "Index: idx_leave_allocation_year_emp",
        "Index: idx_monthly_payroll_period_status",
        "Index: idx_monthly_payroll_emp_period",
        "Index: idx_payroll_component_type_name",
        "Index: idx_disbursement_status_date",
        "Index: idx_audit_table_record",
        "Index: idx_audit_timestamp_user"
    ]
    for idx_name in indexes:
        add_check_item(idx_name)

    add_sub_header("Database Views (6 Views)")
    views = [
        "View: vw_EmployeeDirectory",
        "View: vw_AttendanceMonthlySummary",
        "View: vw_MonthlyPayrollSummary",
        "View: vw_LeaveBalanceAndPendingRequests",
        "View: vw_DepartmentPayrollExpenditure",
        "View: vw_DisbursementReconciliation"
    ]
    for v in views:
        add_check_item(v)

    # 2. Programmable Logic
    add_section_header("2. Programmable Server-Side Database Objects")
    
    add_sub_header("Stored Procedures (6 Procedures)")
    procedures = [
        "Stored Procedure: sp_ProcessMonthlyPayrollBatch",
        "Stored Procedure: sp_FinalizeMonthlyPayroll",
        "Stored Procedure: sp_OnboardEmployeeWithSalary",
        "Stored Procedure: sp_DisbursePayroll",
        "Stored Procedure: sp_ApproveLeaveRequest",
        "Stored Procedure: sp_RejectLeaveRequest"
    ]
    for sp in procedures:
        add_check_item(sp)

    add_sub_header("User-Defined Functions (6 Scalar UDFs)")
    udfs = [
        "Function: fn_CalculateMonthlyNetPay (DETERMINISTIC NO SQL)",
        "Function: fn_GetEmployeeLeaveBalance (READS SQL DATA)",
        "Function: fn_GetDepartmentMonthlyPayrollTotal (READS SQL DATA)",
        "Function: fn_MaskBankAccount (DETERMINISTIC NO SQL)",
        "Function: fn_CalculateOvertimePay (DETERMINISTIC NO SQL)",
        "Function: fn_GetEmployeeTenureMonths (READS SQL DATA)"
    ]
    for udf in udfs:
        add_check_item(udf)

    add_sub_header("Active Integrity & Audit Triggers (4 Triggers)")
    triggers = [
        "Trigger: trg_AuditSalaryModifications (AFTER UPDATE ON Salary_Structure)",
        "Trigger: trg_EnforceLeaveQuota (BEFORE UPDATE ON Leave_Request)",
        "Trigger: trg_SyncLeaveAllocationOnStatusChange (AFTER UPDATE ON Leave_Request)",
        "Trigger: trg_ValidateAttendanceTimes (BEFORE INSERT ON Attendance_Log)"
    ]
    for trg in triggers:
        add_check_item(trg)

    # 3. DML
    add_section_header("3. Data Manipulation Language (DML) Operations")
    
    add_sub_header("INSERT Operations")
    dml_inserts = [
        "Reference Department Catalog Seeding",
        "Designation Catalog & Salary Band Seeding",
        "Statutory Leave Policies Seeding",
        "User Account Credential Ingestion",
        "Employee Profile Data Ingestion",
        "Effective-Dated Salary Structure Ingestion",
        "Biometric Daily Attendance Punch Ingestion",
        "Annual Leave Entitlement Allocation Ingestion",
        "Leave Application Request Ingestion",
        "Monthly Batch Payroll Ledger Generation",
        "Itemized Payroll Component Details Ingestion",
        "Electronic Fund Transfer (BEFTN) Disbursement Voucher Generation",
        "Structured JSON Audit Delta Log Ingestion"
    ]
    for item in dml_inserts:
        add_check_item(item)

    add_sub_header("SELECT & Analytical Query Operations")
    dml_selects = [
        "Keyset & Offset Pagination Query",
        "Parameterized Multi-Attribute Search Query",
        "Recursive CTE Organizational Hierarchy Traversal",
        "Dense Rank & Rank Compensation Distribution",
        "Month-over-Month Departmental Payroll Variance (LAG)",
        "Day-over-Day Consecutive Shift Variance (LEAD)",
        "Chronic Absenteeism & Punctuality Outlier Detection",
        "Overtime Outliers & Departmental Benchmark Comparison",
        "Statutory Tax & Provident Fund Remittance Statement",
        "Designation Salary Band Compliance Audit",
        "BEFTN Electronic Fund Transfer Reconciliation",
        "Corporate Leave Utilization & Outstanding Liability Report",
        "Compensation Quartile & Percentile Grouping (NTILE)",
        "Structured JSON Audit Trail Inspection",
        "Correlated Subquery Above-Average Salary Filter",
        "Multidimensional Rollup Headcount & Cost Aggregations (WITH ROLLUP)",
        "Cumulative Payroll Spend Running Total",
        "Query Plan Execution Analysis (EXPLAIN FORMAT=TREE / EXPLAIN ANALYZE)",
        "Automated Data Integrity & Invariant Health Scanner (UNION ALL)"
    ]
    for item in dml_selects:
        add_check_item(item)

    add_sub_header("UPDATE Operations")
    dml_updates = [
        "Leave Request Status Update: Pending to Approved",
        "Leave Request Status Update: Pending to Rejected",
        "Leave Request Status Update: Approved to Cancelled",
        "Automatic Leave Quota Allocation Decrement",
        "Automatic Leave Quota Allocation Restoration",
        "Monthly Payroll Status Update: Draft to Finalized",
        "Monthly Payroll Status Update: Finalized to Disbursed",
        "Salary Structure Terms Revision",
        "Biometric Attendance Log Timestamp & Hours Correction",
        "Employee Banking Particulars Update",
        "Department Head Re-assignment"
    ]
    for item in dml_updates:
        add_check_item(item)

    add_sub_header("DELETE & TRUNCATE Operations")
    dml_deletes = [
        "Draft Unfinalized Payroll Records Purge",
        "Cascading Employee Profile Data Deletion",
        "Restricted Historical Payroll Deletion Protection",
        "Restricted Disbursed Voucher Deletion Protection",
        "Reference Catalog Table Truncation (TRUNCATE TABLE)",
        "Post-Test QA Fixture Clean-Up"
    ]
    for item in dml_deletes:
        add_check_item(item)

    # 4. TCL
    add_section_header("4. Transaction Control Language (TCL) & Concurrency Operations")
    tcl = [
        "Explicit Transaction Initiation (START TRANSACTION)",
        "Transaction State Committal (COMMIT)",
        "Transaction State Rollback (ROLLBACK)",
        "Procedural Exception Handling (DECLARE EXIT HANDLER FOR SQLEXCEPTION)",
        "Pessimistic Row-Level Locking: Leave Allocation (SELECT ... FOR UPDATE)",
        "Pessimistic Row-Level Locking: Monthly Payroll Voucher (SELECT ... FOR UPDATE)",
        "Lock Wait Timeout Configuration (SET innodb_lock_wait_timeout = 1)",
        "Application-Level Transaction Boundary (BeginTransactionAsync / CommitAsync)"
    ]
    for item in tcl:
        add_check_item(item)

    # 5. DCL
    add_section_header("5. Data Control Language (DCL) & Security Operations")
    dcl = [
        "Service Account Provisioning (paypulse_user, paypulse_api_svc, paypulse_hr_svc, paypulse_payroll_svc)",
        "Enterprise Role Provisioning (role_paypulse_admin, role_paypulse_hr, role_paypulse_payroll, role_paypulse_supervisor, role_paypulse_employee, role_paypulse_api)",
        "Operational DML Privileges Grant (SELECT, INSERT, UPDATE, DELETE)",
        "Stored Routine Privileges Grant (EXECUTE)",
        "Destructive DDL Privilege Revocation / Prohibition (DROP, ALTER, TRUNCATE)",
        "Password Hash Encryption (BCrypt Work Factor 12)",
        "Financial Account Masking (fn_MaskBankAccount)",
        "Parameterized Query Execution (Anti-SQL Injection)"
    ]
    for item in dcl:
        add_check_item(item)

    # 6. ORM
    add_section_header("6. Application Layer & ORM Operations (.NET 10 / EF Core)")
    orm = [
        "EF Core DbContext Entity Mapping (AppDbContext)",
        "Fluent API Table, Column & Precision Mapping (EntityConfigurations.cs)",
        "Transient Connection Drop Auto-Retry Resiliency (EnableRetryOnFailure)",
        "Navigation Property Eager Loading (Include, ThenInclude)",
        "Read-Only Performance Query Optimization (AsNoTracking)",
        "Filtered Entity Pagination Query (Skip, Take, CountAsync)",
        "Change Tracker Unit of Work Batch Persistence (SaveChangesAsync)",
        "Application-Side Audit Event Dispatching (SystemAuditLog)"
    ]
    for item in orm:
        add_check_item(item)

    # 7. Mock DB
    add_section_header("7. Client-Side Simulated Database Operations (mockDb.ts)")
    mock = [
        "LocalStorage Persistent Storage Emulation",
        "Asynchronous Roundtrip Network Delay Simulation (delay)",
        "Standalone Mock Entity CRUD Operations",
        "Synthetic Mock Store Full Database Reset (mockDb.resetAll)"
    ]
    for item in mock:
        add_check_item(item)

    # 8. Testing & Diagnostics
    add_section_header("8. Verification, Testing & Diagnostic Operations")
    testing = [
        "Standalone SQL Functional Verification Suite (10 Test Queries in 08_test_queries.sql)",
        "Automated Pytest Verification Suite (18 Test Cases in test_database.py)",
        "ASP.NET Core Web API Integration Test Suites (xUnit)",
        "Database Health & Anomaly Invariant Scanners"
    ]
    for item in testing:
        add_check_item(item)

    doc.save(output_path)
    print(f"Successfully generated DOCX checklist at {output_path}")

if __name__ == '__main__':
    create_checklist_docx('/home/noir/Projects/PayPulse/ASSIGNMENT/DATABASE_OPERATIONS_CHECKLIST.docx')
